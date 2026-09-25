# -*- coding: utf-8 -*-
"""
사외 연습 체크리스트 → docs/practice_checklist.xlsx
밖에서 LLM 없이 하나씩 돌리며 "확인했다"를 체크하는 용도. 시트 2개뿐.
"""
import os
import re
import glob

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as L

HERE = os.path.dirname(os.path.abspath(__file__))
LAB = os.path.abspath(os.path.join(HERE, '..', '..'))
OUT = os.path.join(HERE, '..', '..', 'docs', 'practice_checklist.xlsx')

FONT, PRI = '맑은 고딕', '185463'
thin = Side(style='thin', color='BFCACB')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
F = lambda c: PatternFill('solid', start_color=c, end_color=c)
HDR, INP, CALC, WARN, OKF, MUT = F('E7EDEE'), F('FFFDE7'), F('EEF4F5'), F('F5EAD0'), F('E2F0D9'), F('F5F5F5')
f = lambda **k: Font(name=FONT, size=k.pop('size', 10), **k)
WRAP = Alignment(wrap_text=True, vertical='top')
CEN = Alignment(horizontal='center', vertical='center', wrap_text=True)

# 노트북 셀 수 자동 집계
CELLS = {}
for p in sorted(glob.glob(os.path.join(LAB, 'nb*.py'))):
    CELLS[os.path.basename(p)[:4]] = len(re.findall(r'^# %%', open(p, encoding='utf-8').read(), re.M))

wb = Workbook()


def head(ws, title, sub, cols, widths, row=4, n_input=0, input_cols=()):
    ws['A1'] = title; ws['A1'].font = f(size=13, bold=True, color=PRI)
    ws['A2'] = sub; ws['A2'].font = f(size=9, color='6B6B6B')
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(cols))
    for j, c in enumerate(cols, 1):
        x = ws.cell(row, j, c); x.font = f(bold=True); x.fill = HDR; x.border = BOX; x.alignment = CEN
        ws.column_dimensions[L(j)].width = widths[j - 1]
    ws.freeze_panes = ws.cell(row + 1, 1)
    ws.page_setup.orientation = 'landscape'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0


def rows(ws, data, start=5, cen=(), inp=(), bold=()):
    for i, r in enumerate(data, start):
        for j, v in enumerate(r, 1):
            x = ws.cell(i, j, v); x.border = BOX; x.font = f(bold=(j in bold))
            x.alignment = CEN if j in cen else WRAP
            if j in inp:
                x.fill = INP
    return start + len(data)


def dv(ws, items, rng):
    d = DataValidation(type='list', formula1='"' + ','.join(items) + '"', allow_blank=True)
    ws.add_data_validation(d); d.add(rng)



# ================= 1. 사외 연습 체크리스트 =================
ws = wb.active; ws.title = '1_사외연습_체크리스트'
cols = ['단계', '무엇을 한다', '확인할 것 (이게 보이면 정상)', '어디서', '확인', '메모']
head(ws, '1. 사외 연습 체크리스트', '밖에서 LLM 없이 하나씩 돌리며 확인 · E · F 열 기입 · 설정 · 키 불필요',
     cols, [22, 40, 58, 22, 8, 26])
D1 = [
 ('준비', '압축 풀고 pandas · numpy · openpyxl 설치', '[0] 준비 셀에 LAB_MODE=sample · SAMPLE_SET=100 이 찍힘', '아무 노트북', '', ''),
 ('준비', '안내서 읽기', '돌릴 노트북 6개와 순서를 안다', 'docs/practice_guide.html', '', ''),
 ('1 · 데이터', 'nb01 실행 — 소스별 분포', '문서 100건 · 뉴스 32 · 증권사 18 … 로 나옴', 'nb01 [2] 셀', '', ''),
 ('1 · 데이터', '본문 길이 · 기간 확인', '2026-04 ~ 09 · 본문 중앙값 400자 내외', 'nb01 [3] 셀', '', ''),
 ('2 · 청킹', 'nb02 실행 — 정제 전후 비교', '메일 서명 · 기사 말머리가 제거되는 것을 눈으로 확인', 'nb02 [5] 셀', '', ''),
 ('2 · 청킹', '청크 길이 분포', '청크 182개 · 중앙 270자 내외 · 토큰 초과율 0', 'nb02 [3] 셀', '', ''),
 ('2 · 청킹', '기준 대비 재청킹 비교', 'lab_sources 의 chunk 값을 바꾸면 청크 수가 달라짐', 'nb02 [+] 셀', '', ''),
 ('3 · 적재점검', 'nb00 실행 — 무결성 점검', '미달 3건(청킹 누락 · 고아 청크 · 초단문)이 검출됨', 'nb00 [8] 셀', '', ''),
 ('3 · 적재점검', '소스 유형 매핑 확인', '데이터의 모든 코드가 정의돼 있음 메시지', 'nb00 [1b] 셀', '', ''),
 ('4 · 토크나이저', 'nb03 실행 — 3종 비교', 'space 보다 josa 가 hit@1 이 높게 나옴', 'nb03 [3] 셀', '', ''),
 ('4 · 토크나이저', '토큰 분해 눈으로 보기', '"하이닉스의" 가 space 에서만 다른 토큰', 'nb03 [1] 셀', '', ''),
 ('5 · 검색지표', 'nb06 실행 — 방식 6종 비교', 'bm25 R@5 0.889 · dense 0.704 처럼 값이 갈림', 'nb06 [3] 셀', '', ''),
 ('5 · 검색지표', '유형별 · 소스별 약점 보기', '비교형 질문이 낮게 나오는 것 확인', 'nb06 [4] 셀', '', ''),
 ('5 · 검색지표', '틀린 문항 직접 보기', '어떤 청크가 잘못 올라왔는지 원문 확인', 'nb06 [6] · [7] 셀', '', ''),
 ('6 · 답변지표', 'nb10 실행 — 방식 2종 비교', '모범답변(상한) vs 추출식 표가 나옴', 'nb10 [4] 셀', '', ''),
 ('6 · 답변지표', '합격 판정 읽기', '충실도 · 인용은 충족 / 정확도는 미달로 나옴', 'nb10 [5] 셀', '', ''),
 ('6 · 답변지표', '약한 문항 확인', '정확도 낮은 5문항의 답변을 직접 읽어봄', 'nb10 [6] 셀', '', ''),
 ('실험 1', '토크나이저 바꾸기', "BM25_TOK 를 'space' → 'josa' 로 바꾸면 점수 상승", 'nb06 [2] 셀', '', ''),
 ('실험 2', '검색 조합 바꾸기', 'MODES 에서 모드를 빼고 더해 차이 확인', 'nb06 [3] 셀', '', ''),
 ('실험 3', '청크 크기 바꾸기', 'lab_sources 의 chunk=500 → 300 후 nb06 재실행', 'lab_sources.py', '', ''),
 ('실험 4', 'RRF k · MMR λ 스윕', '결합 · 다양성 파라미터가 점수에 주는 영향', 'nb06 [5] 셀', '', ''),
 ('실험 5', '정답지 직접 수정', '질문을 바꿔 넣고 재실행 → 약한 질문 유형 파악', 'golden/golden_sample100.csv', '', ''),
 ('정리', '사내에서 고칠 3군데 파악', '.env · lab_config(RAW·CHUNK·MV) · lab_sources(SOURCES)', 'practice_guide 6장', '', ''),
]
last = rows(ws, D1, cen=(5,), inp=(5, 6), bold=(1,)) - 1
dv(ws, ['확인', '보류', '질문 있음'], f'E5:E{last}')
for v, c in (('확인', 'E2F0D9'), ('질문 있음', 'F5EAD0')):
    ws.conditional_formatting.add(f'E5:E{last}', CellIsRule(operator='equal', formula=[f'"{v}"'], fill=F(c)))
b = last + 2
ws.cell(b, 1, '집계').font = f(bold=True, color=PRI)
for i, (k, v) in enumerate([('전체', f'=COUNTIF(B5:B{last},"<>")'),
                            ('확인 완료', f'=COUNTIF(E5:E{last},"확인")'),
                            ('질문 있음', f'=COUNTIF(E5:E{last},"질문 있음")'),
                            ('진행률', f'=IFERROR(COUNTIF(E5:E{last},"확인")/COUNTIF(B5:B{last},"<>"),"")')], b + 1):
    ws.cell(i, 1, k).font = f(bold=True); ws.cell(i, 1).border = BOX
    c = ws.cell(i, 2, v); c.fill = CALC; c.border = BOX; c.font = f()
    if k == '진행률':
        c.number_format = '0%'

# ================= 1b. 연습 순서 요약 =================
ws = wb.create_sheet('1b_노트북_연습순서')
cols = ['순서', '실행 파일', '무엇을 보는가', 'LLM', '소요', '완료', '메모']
head(ws, '1b. 연습 순서 요약', '밖에서 돌리는 6개 · LLM 불필요 · 나머지 8개는 사내에서',
     cols, [7, 26, 58, 9, 9, 8, 26])
D1b = [
 (1, 'nb01_inventory', '데이터가 어떻게 생겼나 — 소스별 분포 · 기간 · 길이', '불필요', '20분', '', ''),
 (2, 'nb02_parse_chunk', '청크로 어떻게 쪼개지나 · 정제 전후', '불필요', '30분', '', ''),
 (3, 'nb00_pipeline_check', '적재 문제를 잡아내는가 (미달 3건)', '불필요', '20분', '', ''),
 (4, 'nb03_bm25_tokenizer', '한국어 조사 처리의 영향', '불필요', '20분', '', ''),
 (5, 'nb06_retrieval_eval', '★ 검색 지표 — 방식 6종 비교', '불필요', '40분', '', ''),
 (6, 'nb10_answer_eval', '★ 답변 지표 — 모범답변 vs 추출식', '불필요', '30분', '', ''),
 ('-', 'nb12 · nb13', '골든셋 생성 · 프리셋 비교 — 사내에서', '필요', '-', '', ''),
 ('-', 'nb04 · nb05 · nb07 · nb08 · nb09 · nb11', 'Milvus · 스모크 · 정형연계 · 자동평가 · 양식 연결', '일부 필요', '-', '', ''),
]
last = rows(ws, D1b, cen=(1, 4, 5, 6), inp=(6, 7), bold=(2,)) - 1
dv(ws, ['완료', '진행', '미착수'], f'F5:F{last}')


from openpyxl.workbook.properties import CalcProperties      # noqa: E402
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.active = 0
wb.save(OUT)
print('saved', os.path.abspath(OUT), [w.title for w in wb.worksheets])
