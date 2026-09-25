# -*- coding: utf-8 -*-
"""
파일 · 함수 목록 → docs/rag_file_index_rev1.xlsx
코드를 직접 파싱해서 만든다 (손으로 적지 않음 — 코드가 바뀌면 다시 돌리면 됨).
  1_파일목록   : rag_lab 안의 모든 파일 · 역할 · 확인 칸
  2_함수목록   : 모듈별 함수 · 인자 · 한 줄 설명 · 줄 번호 · 확인 칸
  3_노트북_셀  : 노트북별 셀 순서 · 셀 제목 · 확인 칸
"""
import ast
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as L

HERE = os.path.dirname(os.path.abspath(__file__))
LAB = os.path.abspath(os.path.join(HERE, '..', '..'))
OUT = os.path.join(LAB, 'docs', 'rag_file_index_rev1.xlsx')

FONT, PRI = '맑은 고딕', '185463'
thin = Side(style='thin', color='BFCACB')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
F = lambda c: PatternFill('solid', start_color=c, end_color=c)
HDR, INP, CALC, WARN, OKF, MUT = F('E7EDEE'), F('FFFDE7'), F('EEF4F5'), F('F5EAD0'), F('E2F0D9'), F('F5F5F5')
f = lambda **k: Font(name=FONT, size=k.pop('size', 10), **k)
WRAP = Alignment(wrap_text=True, vertical='top')
CEN = Alignment(horizontal='center', vertical='center', wrap_text=True)

SKIP_DIR = {'.git', '__pycache__', 'out', '.ipynb_checkpoints'}

# 파일 역할 — 코드에서 뽑을 수 없는 것만 수기
ROLE = {
 'lab_config.py': ('설정', '접속 정보 · 테이블 · 컬럼 매핑 · 실행 모드. 사내에서 [사내 맞춤] 블록 수정'),
 'lab_sources.py': ('설정', '소스 카테고리 정의 — 코드 · 이름 · 사내 DB값(alias) · 필수 메타 · 정제 · 청킹'),
 'lab_io.py': ('공통', 'Oracle 읽기 / 샘플 제공 / 결과 저장 — 표준 컬럼 DataFrame 반환'),
 'lab_text.py': ('공통', '정제 · 토큰 수 · 품질 지표 · 근사 중복 · 유형별 청킹'),
 'lab_search.py': ('공통', 'BM25 · 임베딩 · 리랭커 · Milvus · RRF · MMR · LLM 답변'),
 'lab_pipeline.py': ('공통', '적재 무결성 점검 — 기준(CRITERIA) · 문서/메타/청크/인덱스 점검'),
 'lab_eval.py': ('공통', '검색 지표 — 적중 판정 · Hit · Recall · MRR · nDCG'),
 'lab_ragas.py': ('공통', '표준 지표 16종 — 생성(RAGAS 계열) 계산 · LLM 채점 프롬프트'),
 'lab_autoeval.py': ('공통', '라벨 없는 평가셋 생성 · 근거 문서 연결'),
 'lab_golden_llm.py': ('공통', 'LLM 골든셋 생성 · 자가검증 · 유형 배분 · 검수 시트'),
 'lab_graph.py': ('공통', 'LangGraph RAG 그래프 · 프리셋 v0~v6 · 비교'),
 'lab_sql.py': ('공통', '정형 연계 — 라우팅 · SQL 생성 · 정적 검증 · 실행'),
 'lab_sample.py': ('샘플', '샘플 문서 16건 + 청크 + 골든셋 11 (코드가 만들어 냄 · DB 아님)'),
 'lab_sample100.py': ('샘플', '샘플 문서 100건 + 청크 182 + 골든셋 30 (SAMPLE_SET=100)'),
 'lab_sample_sql.py': ('샘플', '샘플 정형 DB(SQLite) + SQL 골든셋 7'),
 'USAGE.md': ('문서', '상세 사용법 · 카테고리 수정법 · Windows FAQ'),
 'README.md': ('문서', '개요 · 실행 방법 요약'),
 'requirements.txt': ('설정', '필수 pandas · numpy · openpyxl / 나머지는 선택'),
 '.env.example': ('설정', '접속 정보 양식 — 복사해서 .env 로 쓴다'),
 '.gitignore': ('설정', '.env · out · __pycache__ 제외'),
}
NB_ROLE = {
 'nb00': ('핵심', '적재 검증 — 평가 이전 관문'), 'nb01': ('선택', '적재 현황 조사'),
 'nb02': ('핵심', '파싱 · 청킹 품질'), 'nb03': ('선택', 'BM25 토크나이저 비교'),
 'nb04': ('선택', 'Milvus 인덱스 · 건수 정합'), 'nb05': ('선택', '스모크 20문항'),
 'nb06': ('핵심', '검색 지표 측정'), 'nb07': ('선택', '정형 연계'),
 'nb08': ('선택', '라벨 없는 자동 평가'), 'nb09': ('참고', '골든셋 구성 → 평가 연결'),
 'nb10': ('핵심', '답변 품질 측정'), 'nb11': ('선택', '질문 ↔ 근거 연결'),
 'nb12': ('핵심', 'LLM 골든셋 생성'), 'nb13': ('핵심', '프리셋 비교 · 개선 정량화'),
}
DIR_ROLE = {
 'golden': ('데이터', '평가용 골든셋 · 양식 · 예시'),
 'catalog': ('데이터', '정형 지표 정의서'),
 'docs': ('문서', '돌리는 데 필요한 문서 3종'),
 'archive': ('보관', '기입용 엑셀 · 참고 문서 · 생성 스크립트'),
 'archive/src': ('보관', '문서 · 엑셀 생성 스크립트'),
 'out': ('결과', '노트북 실행 결과 (zip 제외)'),
}

wb = Workbook()


def head(ws, title, sub, cols, widths, row=4):
    ws['A1'] = title; ws['A1'].font = f(size=13, bold=True, color=PRI)
    ws['A2'] = sub; ws['A2'].font = f(size=9, color='6B6B6B')
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(cols))
    for j, c in enumerate(cols, 1):
        x = ws.cell(row, j, c); x.font = f(bold=True); x.fill = HDR; x.border = BOX; x.alignment = CEN
        ws.column_dimensions[L(j)].width = widths[j - 1]
    ws.freeze_panes = ws.cell(row + 1, 1)
    ws.page_setup.orientation = 'landscape'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0


def put(ws, data, cen=(), inp=(), bold=(), start=5):
    for i, r in enumerate(data, start):
        for j, v in enumerate(r, 1):
            x = ws.cell(i, j, v); x.border = BOX; x.font = f(bold=(j in bold))
            x.alignment = CEN if j in cen else WRAP
            if j in inp:
                x.fill = INP
    return start + len(data)


def check_col(ws, col, last, items=('확인', '보류', '질문 있음')):
    d = DataValidation(type='list', formula1='"' + ','.join(items) + '"', allow_blank=True)
    ws.add_data_validation(d); d.add(f'{col}5:{col}{last}')
    ws.conditional_formatting.add(f'{col}5:{col}{last}',
                                  CellIsRule(operator='equal', formula=['"확인"'], fill=OKF))
    ws.conditional_formatting.add(f'{col}5:{col}{last}',
                                  CellIsRule(operator='equal', formula=['"질문 있음"'], fill=WARN))


def first_doc(node):
    d = ast.get_docstring(node)
    if not d:
        return ''
    line = [l.strip() for l in d.strip().splitlines() if l.strip()]
    return line[0] if line else ''


def sig(node):
    a = [x.arg for x in node.args.args]
    if node.args.vararg:
        a.append('*' + node.args.vararg.arg)
    if node.args.kwarg:
        a.append('**' + node.args.kwarg.arg)
    return '(' + ', '.join(a) + ')'


# ================= 0. 안내 =================
ws = wb.active; ws.title = '0_안내'
ws['A1'] = '파일 · 함수 목록 rev.1 — 하나씩 보며 확인 체크'; ws['A1'].font = f(size=14, bold=True, color=PRI)
ws['A2'] = '2026-09-25 · 코드를 직접 파싱해 생성 (archive/src/make_file_index.py) · 연노랑 = 기입'
ws['A2'].font = f(size=9, color='6B6B6B')
intro = [
 ('쓰는 법', '1_파일목록 → 2_함수목록 → 3_노트북_셀 순서로 훑으며 "확인" 열 기입'),
 ('확인 열', '확인 / 보류 / 질문 있음 — 질문 있음은 노랑으로 표시됨'),
 ('구분', '설정 = 사내에서 고치는 것 · 공통 = 노트북이 불러 쓰는 코드 · 샘플 = 연습용 가짜 데이터'),
 ('핵심 표시', '노트북 구분의 핵심 6개(nb00 · nb02 · nb06 · nb10 · nb12 · nb13)는 초록'),
 ('다시 만들기', '코드가 바뀌면 archive/src/make_file_index.py 를 다시 실행'),
 ('주의', '함수 설명은 코드의 첫 줄 주석(docstring)을 그대로 가져온 것 — 설명이 빈 칸이면 주석이 없는 함수'),
]
ws['A4'] = '항목'; ws['B4'] = '내용'
for c in ('A4', 'B4'):
    ws[c].font = f(bold=True); ws[c].fill = HDR; ws[c].border = BOX; ws[c].alignment = CEN
for i, (a, b) in enumerate(intro, 5):
    ws.cell(i, 1, a).font = f(bold=True); ws.cell(i, 1).border = BOX; ws.cell(i, 1).alignment = WRAP
    ws.cell(i, 2, b).font = f(); ws.cell(i, 2).border = BOX; ws.cell(i, 2).alignment = WRAP
ws.column_dimensions['A'].width = 18; ws.column_dimensions['B'].width = 104

# ================= 1. 파일 목록 =================
files = []
for root, dirs, names in os.walk(LAB):
    dirs[:] = [d for d in sorted(dirs) if d not in SKIP_DIR]
    rel_dir = os.path.relpath(root, LAB).replace('\\', '/')
    if rel_dir == '.':
        rel_dir = ''
    for n in sorted(names):
        if n.endswith(('.pyc', '.DS_Store')):
            continue
        rel = f'{rel_dir}/{n}' if rel_dir else n
        size = os.path.getsize(os.path.join(root, n))
        if n.startswith('nb') and n.endswith('.py'):
            kind, role = NB_ROLE.get(n[:4], ('노트북', ''))
            role = f'{role} — 실행 대상'
        elif n in ROLE:
            kind, role = ROLE[n]
        else:
            kind, role = DIR_ROLE.get(rel_dir, ('기타', ''))
        files.append((rel, kind, role, f'{size/1024:.1f} KB', '', ''))

ws = wb.create_sheet('1_파일목록')
cols = ['경로', '구분', '역할', '크기', '확인', '메모']
head(ws, '1. 파일 목록', f'전체 {len(files)} 개 · E · F 열 기입', cols, [42, 9, 62, 9, 9, 30])
last = put(ws, files, cen=(2, 4, 5), inp=(5, 6), bold=(1,)) - 1
check_col(ws, 'E', last)
for i in range(5, last + 1):
    v = ws.cell(i, 2).value
    if v == '노트북' or ws.cell(i, 3).value and '실행 대상' in str(ws.cell(i, 3).value):
        for j in (1, 2, 3):
            ws.cell(i, j).fill = F('F2F7F7')
    if v == '설정':
        for j in (1, 2, 3):
            ws.cell(i, j).fill = WARN
b = last + 2
ws.cell(b, 1, '집계').font = f(bold=True, color=PRI)
for i, (k, v) in enumerate([('전체 파일', f'=COUNTIF(A5:A{last},"<>")'),
                            ('확인 완료', f'=COUNTIF(E5:E{last},"확인")'),
                            ('질문 있음', f'=COUNTIF(E5:E{last},"질문 있음")'),
                            ('진행률', f'=IFERROR(COUNTIF(E5:E{last},"확인")/COUNTIF(A5:A{last},"<>"),"")')], b + 1):
    ws.cell(i, 1, k).font = f(bold=True); ws.cell(i, 1).border = BOX
    c = ws.cell(i, 2, v); c.fill = CALC; c.border = BOX; c.font = f()
    if k == '진행률':
        c.number_format = '0%'

# ================= 2. 함수 목록 =================
rows = []
for n in sorted(os.listdir(LAB)):
    if not n.endswith('.py'):
        continue
    src = open(os.path.join(LAB, n), encoding='utf-8').read()
    try:
        tree = ast.parse(src)
    except SyntaxError:
        continue
    kind = NB_ROLE.get(n[:4], ('노트북', ''))[0] if n.startswith('nb') else ROLE.get(n, ('', ''))[0]
    for node in tree.body:
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
            rows.append((n, kind, node.name + sig(node), '함수', first_doc(node), node.lineno, '', ''))
        elif isinstance(node, ast.ClassDef):
            rows.append((n, kind, 'class ' + node.name, '클래스', first_doc(node), node.lineno, '', ''))
            for m in node.body:
                if isinstance(m, (ast.FunctionDef, ast.AsyncFunctionDef)):
                    rows.append((n, kind, f'  {node.name}.{m.name}{sig(m)}', '메서드', first_doc(m), m.lineno, '', ''))
        elif isinstance(node, ast.Assign):
            for t in node.targets:
                if isinstance(t, ast.Name) and t.id.isupper() and len(t.id) > 2 and not t.id.startswith('_'):
                    line = src.splitlines()[node.lineno - 1]
                    cmt = line.split('#', 1)[1].strip() if '#' in line else ''
                    if not cmt:                      # 주석이 없으면 값의 모양을 적는다
                        v = node.value
                        if isinstance(v, ast.Dict):
                            cmt = f'딕셔너리 {len(v.keys)} 항목'
                        elif isinstance(v, (ast.List, ast.Tuple, ast.Set)):
                            cmt = f'목록 {len(v.elts)} 개'
                        elif isinstance(v, ast.Constant):
                            cmt = f'값 = {str(v.value)[:60]}'
                        else:
                            try:
                                cmt = ast.unparse(v)[:70]
                            except Exception:
                                cmt = ''
                    rows.append((n, kind, t.id, '설정값', cmt, node.lineno, '', ''))

ws = wb.create_sheet('2_함수목록')
cols = ['파일', '구분', '함수 · 클래스 · 설정값', '종류', '한 줄 설명 (코드 주석)', '줄', '확인', '메모']
head(ws, '2. 함수 목록', f'전체 {len(rows)} 개 · 코드에서 자동 추출 · G · H 열 기입', cols,
     [22, 8, 46, 8, 62, 6, 9, 26])
last = put(ws, rows, cen=(2, 4, 6, 7), inp=(7, 8)) - 1
check_col(ws, 'G', last)
for i in range(5, last + 1):
    k = ws.cell(i, 4).value
    if k == '클래스':
        for j in range(1, 7):
            ws.cell(i, j).fill = F('DAEAED')
    elif k == '설정값':
        for j in range(1, 7):
            ws.cell(i, j).fill = MUT
b = last + 2
ws.cell(b, 1, '집계').font = f(bold=True, color=PRI)
for i, (k, v) in enumerate([('전체', f'=COUNTIF(C5:C{last},"<>")'),
                            ('확인 완료', f'=COUNTIF(G5:G{last},"확인")'),
                            ('진행률', f'=IFERROR(COUNTIF(G5:G{last},"확인")/COUNTIF(C5:C{last},"<>"),"")')], b + 1):
    ws.cell(i, 1, k).font = f(bold=True); ws.cell(i, 1).border = BOX
    c = ws.cell(i, 2, v); c.fill = CALC; c.border = BOX; c.font = f()
    if k == '진행률':
        c.number_format = '0%'

# ================= 3. 노트북 셀 =================
NOISE_CALL = {'print', 'len', 'str', 'int', 'float', 'list', 'dict', 'set', 'sorted', 'range', 'enumerate',
              'zip', 'round', 'abs', 'min', 'max', 'sum', 'open', 'type', 'isinstance', 'importlib.reload',
              'pd.set_option', 'os.path.join', 'sys.path.insert', 'os.chdir', 'os.path.dirname',
              'os.path.abspath', 'os.path.exists', 'format', 'bool', 'to_string', 'head', 'tail',
              'copy', 'items', 'keys', 'values', 'append', 'fillna', 'reset_index', 'astype',
              'strip', 'split', 'join', 'get', 'to_dict', 'value_counts', 'sort_values', 'round'}


def _call_name(node):
    """호출 표현식에서 이름 추출 — lab_io.load_raw / ret.search 형태로"""
    fn = node.func
    parts = []
    while isinstance(fn, ast.Attribute):
        parts.append(fn.attr)
        fn = fn.value
    if isinstance(fn, ast.Name):
        parts.append(fn.id)
    return '.'.join(reversed(parts)) if parts else ''


def cell_summary(code):
    """셀 코드 → (호출하는 함수, 만드는 변수)"""
    try:
        tree = ast.parse(code)
    except SyntaxError:
        return '', ''
    calls, names = [], []
    for node in ast.walk(tree):
        if isinstance(node, ast.Call):
            nm = _call_name(node)
            if nm and nm not in NOISE_CALL and not nm.endswith('.to_string') and nm not in calls:
                calls.append(nm)
        if isinstance(node, ast.Assign):
            for t in node.targets:
                if isinstance(t, ast.Name) and not t.id.startswith('_') and t.id not in names:
                    names.append(t.id)
    return ' · '.join(calls[:6]), ' · '.join(names[:6])


cells = []
for n in sorted(os.listdir(LAB)):
    if not (n.startswith('nb') and n.endswith('.py')):
        continue
    src = open(os.path.join(LAB, n), encoding='utf-8').read()
    kind, role = NB_ROLE.get(n[:4], ('', ''))
    ref = ''
    m = re.search(r'(체크리스트|기입 엑셀): ([^\n]+)', src)
    if m:
        ref = m.group(0)
    marks = list(re.finditer(r'^# %% (.+)$', src, re.M))
    for k, mm in enumerate(marks):
        title = mm.group(1).strip()
        ln = src[:mm.start()].count('\n') + 1
        end = marks[k + 1].start() if k + 1 < len(marks) else len(src)
        body = src[mm.end():end]
        calls, names = cell_summary(body)
        cells.append((n, kind, title, ln, calls, names, ref if k == 0 else '', '', ''))

ws = wb.create_sheet('3_노트북_셀')
cols = ['노트북', '구분', '셀 제목 (Ctrl+Enter 로 하나씩 실행)', '줄', '이 셀이 호출하는 함수',
        '만들어지는 변수 (Variable Explorer)', '참조 문서', '확인', '메모']
head(ws, '3. 노트북 셀 목록', f'전체 {len(cells)} 셀 · 위에서 아래로 실행 · H · I 열 기입', cols,
     [23, 7, 46, 6, 50, 36, 44, 8, 22])
last = put(ws, cells, cen=(2, 4, 8), inp=(8, 9)) - 1
check_col(ws, 'H', last)
prev = None
for i in range(5, last + 1):
    cur = ws.cell(i, 1).value
    if cur != prev:
        for j in range(1, 10):
            ws.cell(i, j).fill = OKF if ws.cell(i, 2).value == '핵심' else F('F2F7F7')
    prev = cur
b = last + 2
ws.cell(b, 1, '집계').font = f(bold=True, color=PRI)
for i, (k, v) in enumerate([('전체 셀', f'=COUNTIF(C5:C{last},"<>")'),
                            ('확인 완료', f'=COUNTIF(H5:H{last},"확인")'),
                            ('진행률', f'=IFERROR(COUNTIF(H5:H{last},"확인")/COUNTIF(C5:C{last},"<>"),"")')], b + 1):
    ws.cell(i, 1, k).font = f(bold=True); ws.cell(i, 1).border = BOX
    c = ws.cell(i, 2, v); c.fill = CALC; c.border = BOX; c.font = f()
    if k == '진행률':
        c.number_format = '0%'

from openpyxl.workbook.properties import CalcProperties      # noqa: E402
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.active = 0
wb.save(OUT)
print('saved', OUT)
print('파일', len(files), '· 함수/설정값', len(rows), '· 셀', len(cells))
