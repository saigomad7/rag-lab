# -*- coding: utf-8 -*-
"""
사내 반입용 이력 관리 대장 → docs/rag_project_log_rev1.xlsx
진행 이력 · 산출물 · 파이프라인 점검 · 골든셋 · 실험 이력 · 지표 현황 · 이슈 · 다음 할 일
연노랑 = 사내 기입 / 회색 = 자동 계산
"""
import os
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

HERE = os.path.dirname(os.path.abspath(__file__))
LAB = os.path.abspath(os.path.join(HERE, '..', '..'))
sys.path.insert(0, LAB)
os.environ.setdefault('LAB_MODE', 'sample')
_cwd = os.getcwd()
os.chdir(LAB)
import lab_ragas          # noqa: E402
import lab_golden_llm     # noqa: E402
import lab_pipeline       # noqa: E402
import lab_graph          # noqa: E402
os.chdir(_cwd)

OUT = os.path.join(HERE, '..', 'rag_project_log_rev1.xlsx')
FONT, PRI = '맑은 고딕', '185463'
thin = Side(style='thin', color='BFCACB')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
F = lambda c: PatternFill('solid', start_color=c, end_color=c)
HDR, INP, CALC, WARN = F('E7EDEE'), F('FFFDE7'), F('EEF4F5'), F('F5EAD0')
f = lambda **k: Font(name=FONT, size=k.pop('size', 10), **k)
WRAP = Alignment(wrap_text=True, vertical='top')
CEN = Alignment(horizontal='center', vertical='center', wrap_text=True)
N = 200

wb = Workbook()


def head(ws, title, sub, cols, widths, n_input=N, input_cols=(), calc_cols=(), row=4):
    ws['A1'] = title; ws['A1'].font = f(size=13, bold=True, color=PRI)
    ws['A2'] = sub; ws['A2'].font = f(size=9, color='6B6B6B')
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(cols))
    for j, c in enumerate(cols, 1):
        x = ws.cell(row, j, c); x.font = f(bold=True); x.fill = HDR; x.border = BOX; x.alignment = CEN
        ws.column_dimensions[L(j)].width = widths[j - 1]
    for i in range(row + 1, row + 1 + n_input):
        for j in range(1, len(cols) + 1):
            x = ws.cell(i, j); x.border = BOX; x.alignment = WRAP; x.font = f()
            if j in input_cols:
                x.fill = INP
            elif j in calc_cols:
                x.fill = CALC
    ws.freeze_panes = ws.cell(row + 1, 1)
    ws.auto_filter.ref = f'A{row}:{L(len(cols))}{row + max(n_input, 1)}'
    ws.page_setup.orientation = 'landscape'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0


def dv(ws, items, rng):
    d = DataValidation(type='list', formula1='"' + ','.join(items) + '"', allow_blank=True)
    ws.add_data_validation(d); d.add(rng)


def fill(ws, rows, start=5, bold_col=None, cen_cols=(), inp_cols=()):
    for i, r in enumerate(rows, start):
        for j, v in enumerate(r, 1):
            x = ws.cell(i, j, v); x.font = f(bold=(j == bold_col)); x.border = BOX
            x.alignment = CEN if j in cen_cols else WRAP
            if j in inp_cols:
                x.fill = INP
    return start + len(rows)


def agg(ws, row, items, wide=2):
    ws.cell(row, 1, '집계').font = f(size=11, bold=True, color=PRI)
    for i, (k, v, fmt) in enumerate(items, row + 1):
        x = ws.cell(i, 1, k); x.font = f(bold=True); x.border = BOX
        c = ws.cell(i, wide, v); c.fill = CALC; c.border = BOX; c.font = f(); c.number_format = fmt


# ================= 0. 안내 =================
ws = wb.active; ws.title = '0_안내'
ws['A1'] = '사내 RAG 구축 · 평가 이력 대장 rev.1'; ws['A1'].font = f(size=14, bold=True, color=PRI)
ws['A2'] = '2026-09-23 · 마켓 인텔리전스 포털용 RAG — 수집 → 적재 → 청킹 → 임베딩 → 인덱싱 → 골든셋 → 평가 → 고도화'
ws['A2'].font = f(size=9, color='6B6B6B')
flow = [
 ('1', '수집 · 파싱', '외부(뉴스 · SNS · 증권사 · 기관) · 내부(보고서 · 메일 · 회의록) → 원본 적재', '-', '완료'),
 ('2', '통합 · 메타', '분산 테이블 → 단일 문서 테이블 · 소스 유형별 메타 설계(공통 컬럼 · META_EXTRA JSON)', 'rag_checklist_rev10.xlsx `소스별_컬럼정의`', '진행'),
 ('3', '청킹', 'DOC_ID 로 연결되는 청크 테이블 적재', 'nb02_parse_chunk', '진행'),
 ('4', '임베딩 · 인덱싱', 'BGE-M3 → Milvus (dense + sparse · 스칼라 필터)', 'nb04_milvus', '진행'),
 ('5', '적재 검증', '단계별 무결성 점검 — 미달 항목 해소 후 평가 진입', 'nb00_pipeline_check → `3_파이프라인점검`', '미착수'),
 ('6', '골든셋', 'LLM 으로 질의 · 응답 · 근거 생성 → 자가 검증 → 사람 검수 → 확정', 'nb12_golden_llm → `4_골든셋현황`', '미착수'),
 ('7', '정량 평가', '검색(IR) · 생성(RAGAS 계열) 지표 측정 · 합격 판정', 'nb06 · nb10 → `7_지표현황`', '미착수'),
 ('8', '고도화', 'LangGraph 노드 구성 변경 → 프리셋별 지표 비교 → 개선 폭 정량화', 'nb13_langgraph_rag → `5_실험이력` · `6_프리셋비교`', '미착수'),
]
ws['A4'] = '단계'; ws['B4'] = '이름'; ws['C4'] = '내용'; ws['D4'] = '도구 · 연결 시트'; ws['E4'] = '상태'
for c in 'ABCDE':
    ws[f'{c}4'].font = f(bold=True); ws[f'{c}4'].fill = HDR; ws[f'{c}4'].border = BOX; ws[f'{c}4'].alignment = CEN
r = fill(ws, flow, 5, bold_col=2, cen_cols=(1, 5), inp_cols=(5,))
dv(ws, ['완료', '진행', '미착수', '보류'], f'E5:E{r - 1}')
rules = [
 ('기입 원칙', '연노랑 칸만 기입 · 회색은 수식 · 회차마다 행 추가(덮어쓰기 금지)'),
 ('측정 성립 조건', '원본 적재 → 파싱 → 청킹 → 인덱싱 완료 구간만 평가 대상 · 미적재 문항은 수집 과제로 분리'),
 ('실험 규칙', '회차당 설정 1개만 변경 · 평가셋 버전 고정 · 차이가 표본오차보다 클 때만 개선 인정'),
 ('파일 위치', '코드 = rag_lab/ · 문서 = rag_lab/docs/ · 실행 결과 = rag_lab/out/'),
 ('실행 환경', 'Windows + Spyder · # %% 셀 단위 실행 · 결과는 Variable Explorer 에서 DataFrame 확인'),
]
r += 1
ws.cell(r, 1, '운영 규칙').font = f(size=11, bold=True, color=PRI)
for i, (a, b) in enumerate(rules, r + 1):
    x = ws.cell(i, 1, a); x.font = f(bold=True); x.border = BOX; x.alignment = WRAP
    y = ws.cell(i, 2, b); y.font = f(); y.border = BOX; y.alignment = WRAP
    ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=5)
for col, w in (('A', 10), ('B', 16), ('C', 62), ('D', 40), ('E', 10)):
    ws.column_dimensions[col].width = w

# ================= 1. 진행 이력 =================
ws = wb.create_sheet('1_진행이력')
HIST = [
 ('2026-09-18', '#1', '진단', '구축된 RAG 워크플로 적정성 · 수준 진단 요청 (평가 이력 없음)', '1차 진단 · 작업 노트', '완료'),
 ('2026-09-18', '#2', '문서', '정리본 HTML 요청', 'rag_diagnosis_rev1.html', '완료'),
 ('2026-09-18', '#3', '문서', '단계별 체크리스트 · 옵션 제시 요청', 'rag_quest_board_rev1.html (32 퀘스트)', '완료'),
 ('2026-09-18', '#4', '설계', '외부 설계안 대조 → 최적안 도출', 'rag_design_guide_rev1.html', '완료'),
 ('2026-09-18', '#5', '설계', '재검토 · 개선점 R1~R8', 'design_guide_rev2 · quest_board_rev2', '완료'),
 ('2026-09-18', '#6', '문서', 'T2 포맷 체크리스트', 'rag_checklist_rev1.html', '완료'),
 ('2026-09-18', '#7', '범위', 'Phase 1 = RAG 평가 · 개선 / Phase 2 = 마켓 센싱 · RRF 와 MMR 위치 정리', '방향 확정', '완료'),
 ('2026-09-18', '#8', '문서', '파이프라인 단계별 체크리스트(적재 포함) · 데이터프레임 형식', 'rag_stage_checklist_rev1.html (21단계)', '완료'),
 ('2026-09-18', '#9', '문서', '도식 보강 요청', 'stage_checklist_rev2 (SVG 6장)', '완료'),
 ('2026-09-19', '#10', '문서', '설계 가이드 · 퀘스트 보드 rev3', 'design_guide_rev3 · quest_board_rev3', '완료'),
 ('2026-09-19', '#11', '제약', '모바일 · 사내 반입 → 정적 HTML 전환(스크립트 0)', 'stage_rev3 · quest_rev4 · guide_rev4', '완료'),
 ('2026-09-19', '#12', '확인', '수집 현황 확인값 반영 · 스모크 20문항 작성', 'smoke_test_20_rev1.html', '완료'),
 ('2026-09-19', '#13', '문서', '스프레드시트형 체크리스트', 'rag_checklist_sheet_rev1.html', '완료'),
 ('2026-09-19', '#14', '설계', '소스 유형별 처리 구분', 'sheet_rev2 · sources_data.py', '완료'),
 ('2026-09-19', '#15', '설계', '소스별 파싱 · 청킹 판단 기준 도식', 'sheet_rev3', '완료'),
 ('2026-09-19', '#16', '산출', '엑셀 병행 요청', 'rag_checklist_rev3.xlsx', '완료'),
 ('2026-09-19', '#17', '판단', '현 수준 점수 추정 (측정 아님)', '현 40~50 → 개선 후 70~78 추정', '완료'),
 ('2026-09-19', '#18', '코드', 'Spyder 셀 단위 실행 코드 요청', 'rag_lab/ (6 노트북 · 샘플 모드)', '완료'),
 ('2026-09-19', '#19', '환경', '사내 = Windows + Spyder 대응', 'BOM · CRLF · cp949 · 엑셀 잠김 처리', '완료'),
 ('2026-09-20', '#20', '배포', 'GitHub public 업로드', 'github.com/saigomad7/rag-lab', '완료'),
 ('2026-09-20', '#21', '설계', 'Phase 3 정형 데이터 연계 정의', 'S22~S29 · 결합 패턴 A~D · sheet_rev4', '완료'),
 ('2026-09-20', '#22', '코드', '지표 정의서 양식 · SQL 라우터', 'metric_catalog_template · lab_sql · nb07', '완료'),
 ('2026-09-20', '#23', '문서', '혼합 질의 처리 흐름 도식', 'query_flow_example_rev1', '완료'),
 ('2026-09-20', '#24', '문서', '지표 정의서 작성 예시 부록', 'query_flow_example_rev2', '완료'),
 ('2026-09-21', '#25', '배포', '문서 일체 공개 저장소 반영', 'docs/ · docs/src/', '완료'),
 ('2026-09-21', '#26', '설계', '소스별 컬럼 정의서(샘플 포함)', 'columns_data.py · sheet_rev5 · xlsx_rev5', '완료'),
 ('2026-09-21', '#27', '설계', '메타 JSON(META_EXTRA · TAGS · ACL) 저장 위치 반영', 'sheet_rev6 · JSON_스키마 시트', '완료'),
 ('2026-09-21', '#28', '보완', '본문 · 요약 컬럼 누락 지적 반영', 'sheet_rev7 (본문 14 · 청크 2 컬럼 추가)', '완료'),
 ('2026-09-22', '#29', '문서', '문서체 개조식 전환', 'sheet_rev8 · xlsx_rev8', '완료'),
 ('2026-09-22', '#30', '평가', '소스 7종 정리 · 라벨 없는 자동 평가셋 설계', 'lab_autoeval · nb08 · rev9', '완료'),
 ('2026-09-23', '#31', '평가', '골든셋 실물 예시 요청', 'golden_set_example_rev1.xlsx · examples/', '완료'),
 ('2026-09-23', '#32', '평가', '정량 평가 기입 양식', 'golden_eval_pack_rev1.xlsx · nb09', '완료'),
 ('2026-09-23', '#33', '평가', '사외 정보 기반 질의 60문항 · 표준 지표 16종', 'answer_golden_v1.csv · lab_ragas · nb10 · pack_rev2', '완료'),
 ('2026-09-23', '#34', '보완', '골든셋 선행 조건(원본 → 파싱 · 청킹) 지적 반영', 'nb11_evidence_link · pack_rev3 (2d · 2e)', '완료'),
 ('2026-09-23', '#35', '확장', '파이프라인 검증 · LLM 골든셋 생성 · LangGraph 고도화 · 이력 대장',
  'lab_pipeline · nb00 / lab_golden_llm · nb12 / lab_graph · nb13 / 본 대장', '완료'),
 ('2026-09-23', '#36', '보완', '기존 엑셀에 신규 도구 반영 · 수식 캐시 중복 주입 버그(파일 손상) 수정',
  'golden_eval_pack_rev4(1b · 프리셋) · rag_checklist_rev10(단계별_실행도구) · golden_set_example_rev2', '완료'),
]
cols = ['일자', '회차', '구분', '요청 · 결정', '산출물', '상태', '사내 메모']
head(ws, '1. 진행 이력', '착수부터의 요청 · 결정 · 산출물 · 신규 행은 아래에 추가', cols,
     [11, 7, 9, 58, 46, 8, 34], n_input=len(HIST) + 40, input_cols=(7,))
HEND = 4 + len(HIST) + 40
fill(ws, HIST, 5, cen_cols=(1, 2, 3, 6))
dv(ws, ['완료', '진행', '보류'], f'F5:F{4 + len(HIST) + 40}')
dv(ws, ['진단', '설계', '문서', '코드', '평가', '환경', '배포', '보완', '확장', '범위', '확인', '판단', '산출'],
   f'C5:C{4 + len(HIST) + 40}')
agg(ws, HEND + 2, [('총 이력', f'=COUNTIF(B5:B{HEND},"<>")', '0'),
                   ('완료', f'=COUNTIF(F5:F{HEND},"완료")', '0'),
                   ('진행 · 보류', f'=COUNTIF(F5:F{HEND},"진행")+COUNTIF(F5:F{HEND},"보류")', '0')])

# ================= 2. 산출물 목록 =================
ws = wb.create_sheet('2_산출물목록')
ART = [
 ('코드', 'lab_config.py', '설정 · 사내 테이블 · 컬럼 매핑', 'rag_lab/', '사내 맞춤 블록 수정 필요'),
 ('코드', 'lab_io.py', 'Oracle · 샘플 입출력 · 결과 저장', 'rag_lab/', ''),
 ('코드', 'lab_text.py', '정제 · 토큰 · 근사 중복', 'rag_lab/', ''),
 ('코드', 'lab_search.py', 'BM25 · 임베딩 · Milvus · RRF · MMR · 리랭크 · 답변', 'rag_lab/', ''),
 ('코드', 'lab_pipeline.py', '적재 파이프라인 무결성 점검', 'rag_lab/', '신규'),
 ('코드', 'lab_eval.py', '검색 지표(hit · recall · MRR · nDCG)', 'rag_lab/', ''),
 ('코드', 'lab_ragas.py', '표준 지표 16종 · LLM 채점', 'rag_lab/', ''),
 ('코드', 'lab_autoeval.py', '라벨 없는 자동 평가셋 · 근거 문서 연결', 'rag_lab/', ''),
 ('코드', 'lab_golden_llm.py', 'LLM 기반 골든셋 생성 · 자가 검증', 'rag_lab/', '신규'),
 ('코드', 'lab_graph.py', 'LangGraph RAG 그래프 · 프리셋 비교', 'rag_lab/', '신규'),
 ('코드', 'lab_sql.py', '정형 데이터 연계(라우팅 · SQL 생성 · 검증)', 'rag_lab/', ''),
 ('노트북', 'nb00_pipeline_check.py', '적재 검증 — 평가 이전 필수', 'rag_lab/', '신규'),
 ('노트북', 'nb01_inventory.py', '적재 현황 조사', 'rag_lab/', ''),
 ('노트북', 'nb02_parse_chunk.py', '파싱 · 청킹 품질', 'rag_lab/', ''),
 ('노트북', 'nb03_bm25_tokenizer.py', '한국어 토크나이저 비교', 'rag_lab/', ''),
 ('노트북', 'nb04_milvus.py', '인덱스 · 검색 파라미터', 'rag_lab/', ''),
 ('노트북', 'nb05_smoke_test.py', '스모크 20문항', 'rag_lab/', ''),
 ('노트북', 'nb06_retrieval_eval.py', '검색 지표 측정', 'rag_lab/', ''),
 ('노트북', 'nb07_sql_router.py', '정형 연계 점검', 'rag_lab/', ''),
 ('노트북', 'nb08_auto_eval.py', '라벨 없는 자동 평가', 'rag_lab/', ''),
 ('노트북', 'nb09_golden_build.py', '골든셋 구성 → 평가 실행', 'rag_lab/', ''),
 ('노트북', 'nb10_answer_eval.py', '답변 품질 표준 지표', 'rag_lab/', ''),
 ('노트북', 'nb11_evidence_link.py', '질문 ↔ 근거 문서 연결', 'rag_lab/', ''),
 ('노트북', 'nb12_golden_llm.py', 'LLM 골든셋 생성 · 검수', 'rag_lab/', '신규'),
 ('노트북', 'nb13_langgraph_rag.py', 'LangGraph 고도화 · 개선 정량화', 'rag_lab/', '신규'),
 ('엑셀', 'rag_run_guide_rev1.xlsx', '실행 가이드 — 순서 · 통과 기준 · 설정 기입표', 'rag_lab/docs/', '신규'),
 ('문서', 'rag_runbook_t2_rev1.html', 'T2 보고서판 안내서 — 도식 13장 · 개념 설명용', 'rag_lab/docs/', '신규'),
 ('문서', 'rag_runbook_rev1.html', '실행 안내서(원본) — 절차 · FAQ 중심', 'rag_lab/docs/', ''),
 ('문서', 'rag_checklist_sheet_rev9.html', '단계별 체크리스트 · 소스별 컬럼 정의', 'rag_lab/docs/', ''),
 ('문서', 'rag_design_guide_rev4.html', '설계 원본', 'rag_lab/docs/', ''),
 ('문서', 'rag_quest_board_rev4.html', '옵션 사전', 'rag_lab/docs/', ''),
 ('문서', 'query_flow_example_rev2.html', '혼합 질의 흐름 · 지표 정의서 예시', 'rag_lab/docs/', ''),
 ('엑셀', 'rag_project_log_rev1.xlsx', '이력 대장 — 본 파일', 'rag_lab/docs/', '신규'),
 ('엑셀', 'rag_checklist_rev10.xlsx', '체크리스트 · 컬럼 정의 · 단계별 실행 도구(S번호↔노트북)', 'rag_lab/docs/', '신규 시트'),
 ('엑셀', 'golden_eval_pack_rev4.xlsx', '골든셋 · 평가 기입본 (LLM 골든셋 검수 · 프리셋 열 포함)', 'rag_lab/docs/', '신규 시트 1b'),
 ('엑셀', 'golden_set_example_rev2.xlsx', '골든셋 실물 예시 — LLM 생성 · 근거 연결 포함 9종', 'rag_lab/docs/', '신규'),
 ('엑셀', 'metric_catalog_template_rev1.xlsx', '정형 지표 정의서 양식', 'rag_lab/docs/', ''),
 ('데이터', 'golden/answer_golden_v1.csv', '사외 정보 기반 질문 60건', 'rag_lab/golden/', '근거 연결 필요'),
 ('데이터', 'golden/smoke20.csv', '스모크 20문항', 'rag_lab/golden/', ''),
]
cols = ['구분', '파일', '용도', '위치', '비고', '사내 반영', '확인자']
head(ws, '2. 산출물 목록', '사내 반입 대상 · F · G 열 기입', cols, [8, 32, 44, 16, 20, 11, 10],
     n_input=len(ART) + 20, input_cols=(6, 7))
fill(ws, ART, 5, cen_cols=(1,))
dv(ws, ['적용', '검토중', '미적용', '해당없음'], f'F5:F{4 + len(ART) + 20}')

# ================= 3. 파이프라인 점검 =================
ws = wb.create_sheet('3_파이프라인점검')
cols = ['단계', '점검 항목', '측정값', '기준', '판정', '비고', '조치 내용', '담당', '기한', '상태']
head(ws, '3. 파이프라인 적재 점검', 'nb00_pipeline_check 출력(out/nb00_pipeline_report_*.csv) 붙여넣기 → G~J 열 기입',
     cols, [12, 34, 11, 11, 8, 40, 40, 9, 11, 9], input_cols=(7, 8, 9, 10))
dv(ws, ['조치 완료', '조치 중', '미착수', '해당없음'], f'J5:J{4 + N}')
crit = pd.DataFrame(list(lab_pipeline.CRITERIA.items()), columns=['키', '기준값'])
b = 4 + N + 2
ws.cell(b, 1, '기준값 (lab_pipeline.CRITERIA)').font = f(size=11, bold=True, color=PRI)
for i, r0 in enumerate(crit.itertuples(), b + 1):
    ws.cell(i, 1, r0.키).font = f(); ws.cell(i, 1).border = BOX
    c = ws.cell(i, 2, r0.기준값); c.border = BOX; c.font = f(); c.number_format = '0.00'
agg(ws, b + len(crit) + 2, [('점검 항목 수', f'=COUNTIF(B5:B{4 + N},"<>")', '0'),
                            ('미달', f'=COUNTIF(E5:E{4 + N},"미달")', '0'),
                            ('조치 완료', f'=COUNTIF(J5:J{4 + N},"조치 완료")', '0'),
                            ('해소율', f'=IFERROR(COUNTIF(J5:J{4 + N},"조치 완료")/COUNTIF(E5:E{4 + N},"미달"),"")', '0.0%')])

# ================= 4. 골든셋 현황 =================
ws = wb.create_sheet('4_골든셋현황')
cols = ['골든셋', '생성 방식', '생성 건수', '자가검증 통과', '검수 채택', '확정 건수', '근거 확보율', '평가 사용', '갱신일', '비고']
head(ws, '4. 골든셋 현황', 'nb12(LLM 생성) · nb11(근거 연결) · nb08(자동 생성) 결과를 회차별로 기입',
     cols, [22, 12, 10, 12, 10, 10, 11, 10, 11, 34], input_cols=tuple(range(1, 11)))
seed = [
 ('golden_llm_v1 (LLM 생성)', 'nb12', '', '', '', '', '', '', '', '코퍼스 기반 · 근거 포함'),
 ('golden_from_bank (사외 60문항 연결)', 'nb11', 60, '', '', '', '', '', '', '근거 확보분만 평가 대상'),
 ('golden_v1 (자동 평가셋)', 'nb08', '', '', '', '', '', '', '', '라벨 0건 · known-item · 제목 · 합성'),
 ('smoke20 (스모크)', '수기', 20, '-', '-', 20, '-', '', '', '체감 점검용'),
]
fill(ws, seed, 5, cen_cols=(2, 3, 4, 5, 6, 7, 8, 9))
mixrow = 5 + len(seed) + 2
ws.cell(mixrow, 1, '유형 배분 (목표 대비)').font = f(size=11, bold=True, color=PRI)
_t = ws.cell(mixrow, 3, '확정 총 건수 →'); _t.font = f(size=9, color='6B6B6B'); _t.alignment = CEN
_tot = ws.cell(mixrow, 4); _tot.fill = INP; _tot.border = BOX; _tot.font = f(bold=True); _tot.alignment = CEN
TOT = f'$D${mixrow}'
mixcols = ['유형', '목표 비중', '목표 건수(확정 기준)', '실제 건수', '비중', '판정']
for j, c in enumerate(mixcols, 1):
    x = ws.cell(mixrow + 1, j, c); x.font = f(bold=True); x.fill = HDR; x.border = BOX; x.alignment = CEN
MIX = list(lab_golden_llm.TARGET_MIX.items())
for i, (t, w) in enumerate(MIX, mixrow + 2):
    ws.cell(i, 1, t).border = BOX; ws.cell(i, 1).font = f(bold=True)
    c = ws.cell(i, 2, w); c.number_format = '0%'; c.border = BOX; c.font = f(); c.alignment = CEN
    c = ws.cell(i, 3, f'=IF({TOT}="","",ROUND({TOT}*B{i},0))'); c.fill = CALC; c.border = BOX; c.font = f(); c.alignment = CEN
    c = ws.cell(i, 4); c.fill = INP; c.border = BOX; c.font = f(); c.alignment = CEN
    c = ws.cell(i, 5, f'=IFERROR(D{i}/SUM($D${mixrow + 2}:$D${mixrow + 1 + len(MIX)}),"")')
    c.fill = CALC; c.border = BOX; c.font = f(); c.number_format = '0.0%'; c.alignment = CEN
    c = ws.cell(i, 6, f'=IF(D{i}="","",IF(D{i}<C{i}*0.6,"부족",IF(D{i}>C{i}*1.5,"과다","적정")))')
    c.fill = CALC; c.border = BOX; c.font = f(); c.alignment = CEN

# ================= 5. 실험 이력 =================
ws = wb.create_sheet('5_실험이력')
cols = ['회차', '일자', '평가셋', '문항 수', '변경 내용 (1회 1개)', '프리셋', 'Recall@5', 'MRR',
        'Faithfulness', 'Answer Correctness', 'Citation', '거절 정확도', 'p95 초', '판정', '비고']
head(ws, '5. 실험 이력', '설정을 하나만 바꾸고 1회차 기록 · 이전 회차 대비 개선 여부 자동 판정',
     cols, [6, 11, 20, 8, 40, 13, 10, 8, 12, 16, 9, 11, 8, 9, 30],
     input_cols=(1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 15), calc_cols=(14,))
for i in range(5, 5 + N):
    c = ws.cell(i, 14, f'=IF(OR(G{i}="",G{i - 1}=""),"",IF(G{i}-G{i - 1}>0.03,"개선",'
                       f'IF(G{i}-G{i - 1}<-0.03,"퇴행","차이 없음")))')
    c.fill = CALC; c.alignment = CEN; c.font = f()
ws['N5'] = ''
dv(ws, list(lab_graph.PRESETS.keys()), f'F5:F{4 + N}')
agg(ws, 4 + N + 2, [('기록 회차', f'=COUNTIF(A5:A{4 + N},"<>")', '0'),
                    ('최고 Recall@5', f'=IFERROR(MAX(G5:G{4 + N}),"")', '0.000'),
                    ('최고 Faithfulness', f'=IFERROR(MAX(I5:I{4 + N}),"")', '0.000'),
                    ('개선 회차', f'=COUNTIF(N5:N{4 + N},"개선")', '0'),
                    ('퇴행 회차', f'=COUNTIF(N5:N{4 + N},"퇴행")', '0')])

# ================= 6. 프리셋 비교 =================
ws = wb.create_sheet('6_프리셋비교')
cols = ['프리셋', '구성', 'Recall@5', 'Faithfulness', 'Answer Relevancy', 'Answer Correctness',
        'Citation', '거절 정확도', '재검색', '재생성', '평균 초', '기준 대비 증감', '채택']
head(ws, '6. 프리셋 비교 — 노드 구성별 성능', 'nb13_langgraph_rag 출력(out/nb13_preset_summary_*.csv) 붙여넣기 · 기준 = v0_baseline',
     cols, [14, 46, 10, 13, 16, 18, 9, 11, 8, 8, 9, 14, 8], n_input=0)
desc = {
 'v0_baseline': 'BM25 단독 — 기준선',
 'v1_hybrid': 'BM25 + 시맨틱 하이브리드(RRF 결합)',
 'v2_rerank': '하이브리드 + 크로스인코더 리랭크',
 'v3_mmr': '리랭크 + MMR 다양성 · 문서당 상한',
 'v4_decompose': '질의 분해(복합 질문 → 하위 질의) 추가',
 'v5_grade': '문단 적합성 판정 + 근거 없을 때 질의 재작성 · 재검색',
 'v6_verify': '답변 환각 점검 → 미달 시 재생성 · 근거 불충분 시 보류',
}
rows = [(k, desc.get(k, ''), '', '', '', '', '', '', '', '', '', '', '') for k in lab_graph.PRESETS]
last = fill(ws, rows, 5, bold_col=1, inp_cols=tuple(range(3, 12)) + (13,)) - 1
for i in range(5, last + 1):
    c = ws.cell(i, 12, f'=IF(OR(C{i}="",$C$5=""),"",ROUND(C{i}-$C$5,3))')
    c.fill = CALC; c.alignment = CEN; c.font = f(); c.number_format = '0.000'
dv(ws, ['운영 적용', '후보', '미채택'], f'M5:M{last}')
note = ['측정 조건 고정: 같은 골든셋 · 같은 코퍼스 버전 · temperature 0',
        '지연이 크게 늘면(평균 초) 노드 이득과 비교 후 채택 결정',
        '채택 프리셋은 lab_graph.DEFAULT 에 반영']
for i, t in enumerate(note, last + 2):
    x = ws.cell(i, 1, t); x.font = f(size=9, color='5C440C'); x.fill = WARN; x.border = BOX
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=len(cols))

# ================= 7. 지표 현황 =================
ws = wb.create_sheet('7_지표현황')
cat = lab_ragas.catalog()
cols = ['구분', '지표', '정의', '계산 방식', '합격 기준', '대응 단계', '1차 측정', '최근 측정', '판정', '비고']
head(ws, '7. 표준 지표 현황', '검색(IR) · 생성(RAGAS 계열) · 운영 16종 · G~H · J 열 기입 (nb06 · nb10 · nb13 결과)',
     cols, [7, 19, 42, 13, 23, 17, 10, 10, 9, 26], n_input=0)
rows = [(r.구분, r.지표, r.정의, r._4, r._5, r._6, '', '', '', '') for r in cat.itertuples()]
last = fill(ws, rows, 5, bold_col=2, cen_cols=(1, 4, 7, 8, 9), inp_cols=(7, 8, 10)) - 1
for i in range(5, last + 1):
    c = ws.cell(i, 9, f'=IF(H{i}="","","기준 대조")'); c.fill = CALC; c.alignment = CEN; c.font = f()
agg(ws, last + 2, [('측정 완료 지표', f'=COUNTIF(H5:H{last},"<>")', '0'),
                   ('전체 지표', len(cat), '0'),
                   ('측정률', f'=IFERROR(COUNTIF(H5:H{last},"<>")/{len(cat)},"")', '0.0%')])

# ================= 8. 이슈 · 결정 =================
ws = wb.create_sheet('8_이슈결정')
ISSUE = [
 ('2026-09-19', '카톡 뷰어에서 HTML 화면이 비어 보임', 'JS 미실행 환경에서 본문을 스크립트로 렌더', '모든 산출 HTML 정적화(스크립트 0)', '완료'),
 ('2026-09-19', '사내 환경 불일치', '맥 기준 코드 작성', 'Windows + Spyder 전제로 전면 수정(BOM · cp949 · 엑셀 잠김)', '완료'),
 ('2026-09-22', '컬럼 정의서에 본문 · 요약 누락', '메타 중심 설계', '본문 테이블 14 · 청크 2 컬럼 추가', '완료'),
 ('2026-09-23', '마케팅 부서 골든셋 확보 곤란', '내부 라벨 작업 부담', '사외 정보 기반 질문 60건 + LLM 생성 골든셋 병행', '완료'),
 ('2026-09-23', '질문만 있고 근거 문서 없음', '질문 우선 작성', 'nb11 근거 연결 단계 신설 · 수집 공백과 검색 실패 분리', '완료'),
 ('', '', '', '', ''),
]
cols = ['일자', '이슈', '원인', '결정 · 조치', '상태', '후속 확인']
head(ws, '8. 이슈 · 결정 사항', '판단 근거를 남겨 재논의 방지 · 신규 행은 아래에 추가', cols,
     [11, 44, 38, 52, 9, 26], n_input=len(ISSUE) + 30, input_cols=(6,))
fill(ws, ISSUE, 5, cen_cols=(1, 5))
dv(ws, ['완료', '진행', '보류'], f'E5:E{4 + len(ISSUE) + 30}')

# ================= 9. 다음 할 일 =================
ws = wb.create_sheet('9_다음할일')
TODO = [
 ('P0', '사내 테이블 · 컬럼 이름을 lab_config [사내 맞춤] 에 반영', '', '', '미착수', 'RAW · CHUNK · MV · DOC_TYPE_MAP'),
 ('P0', '.env 작성 후 LAB_MODE=live 전환', '', '', '미착수', 'Oracle 읽기 전용 계정 사용'),
 ('P0', 'nb00 적재 검증 실행 → 미달 항목 조치', '', '', '미착수', '`3_파이프라인점검` 기입'),
 ('P1', 'nb12 로 LLM 골든셋 생성 → 검수 → 확정', '', '', '미착수', '유형 배분 목표 준수'),
 ('P1', 'nb11 로 사외 60문항 근거 연결 · 수집 공백 목록화', '', '', '미착수', '`2d` · `2e` 시트'),
 ('P1', 'nb06 · nb10 1차 측정 → `7_지표현황` 기입', '', '', '미착수', '기준선 확보'),
 ('P2', 'nb13 프리셋 비교 → 운영 프리셋 확정', '', '', '미착수', '회차당 변경 1개'),
 ('P2', '소스별 컬럼 정의서 사내 보유 여부 기입', '', '', '미착수', 'rag_checklist_rev10.xlsx'),
 ('P2', '정형 지표 정의서 작성 → nb07 연계', '', '', '미착수', 'metric_catalog_template_rev1.xlsx'),
 ('P3', 'Phase 2 마켓 센싱 축별 상시 질의 정의', '', '', '미착수', '6축 · 백테스트'),
]
cols = ['우선', '과제', '담당', '기한', '상태', '비고']
head(ws, '9. 다음 할 일', '우선순위 P0 → P3 · 담당 · 기한 기입', cols, [7, 58, 10, 11, 10, 40],
     n_input=len(TODO) + 30, input_cols=(3, 4, 5, 6))
TEND = 4 + len(TODO) + 30
fill(ws, TODO, 5, cen_cols=(1, 3, 4, 5))
dv(ws, ['완료', '진행', '미착수', '보류'], f'E5:E{4 + len(TODO) + 30}')
dv(ws, ['P0', 'P1', 'P2', 'P3'], f'A5:A{4 + len(TODO) + 30}')
agg(ws, TEND + 2, [('전체 과제', f'=COUNTIF(B5:B{TEND},"<>")', '0'),
                   ('완료', f'=COUNTIF(E5:E{TEND},"완료")', '0'),
                   ('진행률', f'=IFERROR(COUNTIF(E5:E{TEND},"완료")/COUNTIF(B5:B{TEND},"<>"),"")', '0.0%')])

from openpyxl.workbook.properties import CalcProperties      # noqa: E402
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.active = 0
wb.save(OUT)
print('saved', os.path.abspath(OUT), [w.title for w in wb.worksheets])
