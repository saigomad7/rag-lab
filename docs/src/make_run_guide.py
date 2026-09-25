# -*- coding: utf-8 -*-
"""
실행 가이드 → docs/rag_run_guide_rev1.xlsx
무엇을 · 어떤 순서로 돌리고 · 결과를 어디에 적는가. 체크하며 진행하는 용도.
"""
import os
import re
import glob

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as L

HERE = os.path.dirname(os.path.abspath(__file__))
LAB = os.path.abspath(os.path.join(HERE, '..', '..'))
OUT = os.path.join(HERE, '..', 'rag_run_guide_rev1.xlsx')

FONT, PRI = '맑은 고딕', '185463'
thin = Side(style='thin', color='BFCACB')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
F = lambda c: PatternFill('solid', start_color=c, end_color=c)
HDR, INP, CALC, WARN, OKF, MUT = F('E7EDEE'), F('FFFDE7'), F('EEF4F5'), F('F5EAD0'), F('E2F0D9'), F('F5F5F5')
f = lambda **k: Font(name=FONT, size=k.pop('size', 10), **k)
WRAP = Alignment(wrap_text=True, vertical='top')
CEN = Alignment(horizontal='center', vertical='center', wrap_text=True)

# 노트북 셀 수 자동 집계
CELLS = {}
for p in sorted(glob.glob(os.path.join(LAB, 'nb*.py'))):
    CELLS[os.path.basename(p)[:4]] = len(re.findall(r'^# %%', open(p, encoding='utf-8').read(), re.M))

wb = Workbook()


def head(ws, title, sub, cols, widths, row=4, n_input=0, input_cols=()):
    ws['A1'] = title; ws['A1'].font = f(size=13, bold=True, color=PRI)
    ws['A2'] = sub; ws['A2'].font = f(size=9, color='6B6B6B')
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(cols))
    for j, c in enumerate(cols, 1):
        x = ws.cell(row, j, c); x.font = f(bold=True); x.fill = HDR; x.border = BOX; x.alignment = CEN
        ws.column_dimensions[L(j)].width = widths[j - 1]
    ws.freeze_panes = ws.cell(row + 1, 1)
    ws.page_setup.orientation = 'landscape'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0


def rows(ws, data, start=5, cen=(), inp=(), bold=()):
    for i, r in enumerate(data, start):
        for j, v in enumerate(r, 1):
            x = ws.cell(i, j, v); x.border = BOX; x.font = f(bold=(j in bold))
            x.alignment = CEN if j in cen else WRAP
            if j in inp:
                x.fill = INP
    return start + len(data)


def dv(ws, items, rng):
    d = DataValidation(type='list', formula1='"' + ','.join(items) + '"', allow_blank=True)
    ws.add_data_validation(d); d.add(rng)


# ================= 0. 안내 =================
ws = wb.active; ws.title = '0_안내'
ws['A1'] = '실행 가이드 rev.1 — 무엇을 어떤 순서로 돌리는가'; ws['A1'].font = f(size=14, bold=True, color=PRI)
ws['A2'] = '2026-09-25 · 노트북 14개 중 핵심은 6개 · 나머지는 필요할 때만 · 연노랑 = 기입'
ws['A2'].font = f(size=9, color='6B6B6B')
intro = [
 ('왜 14개인가', '한 노트북 = 한 가지 질문. 섞어 놓으면 어느 단계가 문제인지 구분되지 않아 나눠 둠. 전부 돌릴 필요는 없음'),
 ('핵심 6개', 'nb00 적재 검증 · nb12 골든셋 생성 · nb06 검색 측정 · nb10 답변 측정 · nb13 개선 비교 · nb02 청킹 품질'),
 ('선택 8개', 'nb01 현황 · nb03 토크나이저 · nb04 Milvus · nb05 스모크 · nb07 정형연계 · nb08 자동평가 · nb09 골든셋양식 · nb11 근거연결'),
 ('실행 방법', 'Spyder 에서 파일 열고 Ctrl+Enter 로 셀 하나씩. 반드시 [0] 준비 셀부터'),
 ('결과 확인', '콘솔 출력 + Variable Explorer 의 DataFrame. 저장물은 out/ 폴더'),
 ('기록', 'out/ 결과를 이력대장 · 평가팩 엑셀에 붙여넣기. 어느 시트인지는 각 시트의 "기록 위치" 열 참조'),
 ('순서 원칙', '앞 단계 미달 상태에서 뒤 단계를 재면 원인을 특정할 수 없음. 1번부터 차례로'),
]
ws['A4'] = '항목'; ws['B4'] = '내용'
for c in ('A4', 'B4'):
    ws[c].font = f(bold=True); ws[c].fill = HDR; ws[c].border = BOX; ws[c].alignment = CEN
for i, (a, b) in enumerate(intro, 5):
    ws.cell(i, 1, a).font = f(bold=True); ws.cell(i, 1).border = BOX; ws.cell(i, 1).alignment = WRAP
    ws.cell(i, 2, b).font = f(); ws.cell(i, 2).border = BOX; ws.cell(i, 2).alignment = WRAP
r0 = 5 + len(intro) + 1
ws.cell(r0, 1, '시트 안내').font = f(size=11, bold=True, color=PRI)
sheets = [
 ('1_노트북_연습순서', '개인 노트북에서 개념 잡는 순서 (샘플 데이터 100건)'),
 ('2_사내_실행순서', '사내에서 실제로 도는 순서 · 통과 기준 · 기록 위치'),
 ('3_노트북_상세', '14개 각각 무엇을 하는지 · 셀 수 · 입력 · 출력 · 필수 여부'),
 ('4_설정_기입', '.env 와 lab_config 에서 채울 항목 (여기 적고 옮겨 넣기)'),
 ('5_안돌리는_파일', '라이브러리 13개 역할 — 직접 실행하지 않음'),
]
for i, (a, b) in enumerate(sheets, r0 + 1):
    ws.cell(i, 1, a).font = f(bold=True); ws.cell(i, 1).border = BOX
    ws.cell(i, 2, b).font = f(); ws.cell(i, 2).border = BOX; ws.cell(i, 2).alignment = WRAP
ws.column_dimensions['A'].width = 22
ws.column_dimensions['B'].width = 110

# ================= 1. 노트북 연습 순서 =================
ws = wb.create_sheet('1_노트북_연습순서')
cols = ['순서', '실행 파일', '무엇을 보는가', 'LLM 필요', '예상 소요', '완료', '메모']
head(ws, '1. 개인 노트북 연습 순서', '.env → LAB_MODE=sample · SAMPLE_SET=100 · 5~7회차면 개념 정리됨 · F · G 열 기입',
     cols, [7, 26, 62, 10, 11, 8, 30])
D1 = [
 (1, 'nb01_inventory', '문서 100건이 어떻게 생겼나 — 소스별 분포 · 기간 · 본문 길이', '불필요', '20분', '', ''),
 (2, 'nb02_parse_chunk', '청크로 어떻게 쪼개지나 · 정제 전후 차이 · 노이즈 잔존', '불필요', '30분', '', ''),
 (3, 'nb00_pipeline_check', '적재 미달이 어떻게 표시되나 — 일부러 심어둔 결함 3종 검출', '불필요', '20분', '', ''),
 (4, 'nb03_bm25_tokenizer', '조사 처리 방식에 따라 적중률이 달라지는 것', '불필요', '20분', '', ''),
 (5, 'nb06_retrieval_eval', '★ 검색 모드 6종 점수표 — 하이브리드 · 리랭크가 무엇을 바꾸는지', '불필요', '40분', '', ''),
 (6, 'nb08_auto_eval', '라벨 없이 평가셋이 자동 생성되는 과정', '있으면 좋음', '30분', '', ''),
 (7, 'nb12_golden_llm', '★ LLM 이 질문 · 정답 · 근거를 만드는 것 · 자가검증 6항목', '필요', '40분', '', ''),
 (8, 'nb10_answer_eval', '답변이 채점되는 방식 · 충실도 낮은 문항 직접 비교', '필요', '40분', '', ''),
 (9, 'nb13_langgraph_rag', '★ 프리셋 v0~v6 비교표 — 보고에 쓸 결과물 형태', '필요', '40분', '', ''),
]
last = rows(ws, D1, cen=(1, 4, 5, 6), inp=(6, 7), bold=(2,))
dv(ws, ['완료', '진행', '미착수'], f'F5:F{last - 1}')
for v, c in (('완료', 'E2F0D9'), ('진행', 'FFF2CC')):
    ws.conditional_formatting.add(f'F5:F{last - 1}', CellIsRule(operator='equal', formula=[f'"{v}"'], fill=F(c)))
note = [('7 · 8 · 9 는 LLM 필요', '.env 에 SAMPLE_MODELS=true + EMBED_MODE=api + LLM_URL · LLM_API_KEY 기입 (OpenAI 등 외부 API 가능 — 샘플 문서는 지어낸 예시)'),
        ('건너뛰어도 되는 것', 'nb04(Milvus 없음) · nb05(체감용) · nb07(정형 연계는 나중) · nb09 · nb11(엑셀 양식 연결용)'),
        ('확인 포인트', '같은 질문을 검색 모드별로 돌렸을 때 순위가 어떻게 달라지는지 · 왜 그런지 청크 내용을 직접 읽어볼 것')]
b = last + 1
for i, (a, t) in enumerate(note, b):
    x = ws.cell(i, 1, a); x.font = f(bold=True, color='5C440C'); x.fill = WARN; x.border = BOX; x.alignment = WRAP
    y = ws.cell(i, 2, t); y.font = f(color='5C440C'); y.fill = WARN; y.border = BOX; y.alignment = WRAP
    ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=7)

# ================= 2. 사내 실행 순서 =================
ws = wb.create_sheet('2_사내_실행순서')
cols = ['순서', '실행 파일', '목적', '통과 기준 · 판단', '결과 기록 위치', '실행일', '결과', '상태', '메모']
head(ws, '2. 사내 실행 순서', '.env → LAB_MODE=live (SAMPLE_SET · SAMPLE_MODELS 삭제) · F~I 열 기입',
     cols, [7, 26, 40, 40, 34, 11, 26, 9, 26])
D2 = [
 (0, 'lab_config.py 수정', '[사내 맞춤] 블록에 실제 테이블 · 컬럼명 기입', '샘플 없이 load_raw() 가 돌면 성공', '-', '', '', '', ''),
 (1, 'nb00_pipeline_check', '적재 · 메타 · 청킹 · 인덱싱 무결성 점검', '미달 0 건 — 아니면 여기서 멈추고 데이터부터 조치', '이력대장 3_파이프라인점검', '', '', '', ''),
 (2, 'nb01_inventory', '소스별 문서 수 · 기간 · 편중 확인', '수집 공백 기간 파악', '이력대장 3_파이프라인점검', '', '', '', ''),
 (3, 'nb02_parse_chunk', '파싱 · 청킹 품질 — 노이즈 · 표 손상 · 절단', '초단문 비율 · 노이즈 잔존율 확인', '이력대장 3_파이프라인점검', '', '', '', ''),
 (4, 'nb04_milvus', '인덱스 구성 · Oracle 과 건수 일치', '인덱싱 / 청크 ≥ 0.99', '이력대장 3_파이프라인점검', '', '', '', ''),
 (5, 'nb12_golden_llm', '★ 코퍼스에서 골든셋 생성 → 검수 → 확정', '자가검증 통과 ≥ 70% · 유형 배분 적정', '평가팩 1b_LLM골든셋_검수', '', '', '', ''),
 (6, 'nb06_retrieval_eval', '★ 검색 지표 기준선 측정', 'R@5 ≥ 0.80 · R@10 ≥ 0.90 · MRR ≥ 0.65', '평가팩 3_평가_실행기록', '', '', '', ''),
 (7, 'nb10_answer_eval', '★ 답변 지표 기준선 측정', '충실도 ≥ 0.85 · 정확도 ≥ 0.75 · 인용 ≥ 0.90', '이력대장 7_지표현황', '', '', '', ''),
 (8, 'nb13_langgraph_rag', '★ 프리셋별 비교 → 운영 구성 확정', '기준선 대비 증가분 > 표본오차', '이력대장 5_실험이력 · 6_프리셋비교', '', '', '', ''),
 (9, 'nb03_bm25_tokenizer', '한국어 토크나이저 결정', 'Kiwi 적용 시 적중률 개선 폭 확인', '이력대장 5_실험이력', '', '', '', ''),
 (10, 'nb08_auto_eval', '라벨 없는 회귀 점검 (월 1회)', '이전 회차 대비 퇴행 없음', '이력대장 5_실험이력', '', '', '', ''),
 (11, 'nb11_evidence_link', '사외 60문항에 근거 연결 · 수집 공백 목록화', '근거 확보분만 점수 집계', '평가팩 2d · 2e', '', '', '', ''),
 (12, 'nb05_smoke_test', '실제 질문 20건 체감 점검', '실패 유형 5종 집계', '평가팩 스모크 시트', '', '', '', ''),
 (13, 'nb07_sql_router', '정형 데이터 연계 (지표 정의서 작성 후)', '차단 규칙 100% · 실행 정확도 확인', '이력대장 7_지표현황', '', '', '', ''),
 (14, 'nb09_golden_build', '엑셀 양식과 골든셋 연결 (필요 시)', '-', '평가팩 1 · 2', '', '', '', ''),
]
last = rows(ws, D2, cen=(1, 6, 8), inp=(6, 7, 8, 9), bold=(2,))
dv(ws, ['완료', '진행', '미착수', '해당없음'], f'H5:H{last - 1}')
for v, c in (('완료', 'E2F0D9'), ('진행', 'FFF2CC'), ('해당없음', 'F5F5F5')):
    ws.conditional_formatting.add(f'H5:H{last - 1}', CellIsRule(operator='equal', formula=[f'"{v}"'], fill=F(c)))
b = last + 1
ws.cell(b, 1, '집계').font = f(bold=True, color=PRI)
for i, (k, v) in enumerate([('전체', f'=COUNTIF(B5:B{last - 1},"<>")'),
                            ('완료', f'=COUNTIF(H5:H{last - 1},"완료")'),
                            ('진행률', f'=IFERROR(COUNTIF(H5:H{last - 1},"완료")/COUNTIF(B5:B{last - 1},"<>"),"")')], b + 1):
    ws.cell(i, 1, k).font = f(bold=True); ws.cell(i, 1).border = BOX
    c = ws.cell(i, 2, v); c.fill = CALC; c.border = BOX; c.font = f()
    if k == '진행률':
        c.number_format = '0%'
m = b + 5
msg = [('1~4 는 데이터 점검', '여기서 미달이 나오면 5번 이후를 진행해도 점수 해석이 불가능'),
       ('5 는 채점표 만들기', 'LLM 필요. 사내 LLM 엔드포인트가 없으면 nb08 자동 평가셋으로 대체'),
       ('6 · 7 이 기준선', '이 값이 이후 모든 비교의 원점. 설정을 바꾸기 전에 반드시 먼저 측정'),
       ('8 부터 개선', '회차당 설정 1개만 변경 · 같은 골든셋 고정')]
for i, (a, t) in enumerate(msg, m):
    x = ws.cell(i, 1, a); x.font = f(bold=True, color='185463'); x.fill = F('DAEAED'); x.border = BOX; x.alignment = WRAP
    y = ws.cell(i, 2, t); y.font = f(); y.fill = F('DAEAED'); y.border = BOX; y.alignment = WRAP
    ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=9)

# ================= 3. 노트북 상세 =================
ws = wb.create_sheet('3_노트북_상세')
cols = ['파일', '한 줄 목적', '구분', '셀', '입력(무엇을 읽나)', '출력(무엇이 나오나)', '주요 확인 변수', '기록 위치']
head(ws, '3. 노트북 14개 상세', '구분 — 핵심: 반드시 / 선택: 필요할 때 / 참고: 이해용', cols,
     [24, 40, 8, 6, 30, 40, 26, 30])
D3 = [
 ('nb00_pipeline_check', '적재 · 메타 · 청킹 · 인덱싱 무결성 점검', '핵심', CELLS.get('nb00'), '문서 · 청크 테이블',
  '단계별 점검표 · 미달 항목 목록 · 소스별 분포', 'rep · fail · dist', '이력대장 3_파이프라인점검'),
 ('nb01_inventory', '적재 현황 조사 — 소스별 분포 · 기간', '선택', CELLS.get('nb01'), '문서 테이블',
  '소스별 건수 · 기간 · 메타 채움률 · 중복', 'inv · dup', '이력대장 3_파이프라인점검'),
 ('nb02_parse_chunk', '파싱 · 청킹 품질 점검', '핵심', CELLS.get('nb02'), '문서 · 청크',
  '청크 길이 분포 · 노이즈 잔존 · 표 손상 · 문장 절단', 'prof · issues', '이력대장 3_파이프라인점검'),
 ('nb03_bm25_tokenizer', '한국어 토크나이저 3종 비교', '선택', CELLS.get('nb03'), '청크 · 골든셋',
  '공백 / 조사제거 / Kiwi 별 적중률', 'cmp', '이력대장 5_실험이력'),
 ('nb04_milvus', '인덱스 구성 · 건수 정합성', '선택', CELLS.get('nb04'), 'Milvus · 청크',
  '스키마 · 필드 · 건수 비교 · 필터 검색', 'info · cnt', '이력대장 3_파이프라인점검'),
 ('nb05_smoke_test', '실제 질문 20건 체감 점검', '선택', CELLS.get('nb05'), 'golden/smoke20.csv',
  '문항별 검색 결과 · 판정용 엑셀', 'smoke', '평가팩 스모크 시트'),
 ('nb06_retrieval_eval', '검색 지표 측정 — 모드 6종 비교', '핵심', CELLS.get('nb06'), '골든셋 · 청크',
  'Recall@k · MRR · nDCG 모드별 표 · 유형별 분해', 'summary · detail', '평가팩 3_평가_실행기록'),
 ('nb07_sql_router', '정형 데이터 연계 — 라우팅 · SQL · 검증', '선택', CELLS.get('nb07'), '지표 정의서 · SQL 골든셋',
  '라우팅 정확도 · 차단 규칙 · 실행 정확도(EX)', 'route · ex', '이력대장 7_지표현황'),
 ('nb08_auto_eval', '라벨 없는 평가셋 자동 생성 · 측정', '선택', CELLS.get('nb08'), '문서 · 청크',
  'known-item · 제목 · 중복쌍 · 합성 QA 및 기준 대비 판정', 'sets · res', '이력대장 5_실험이력'),
 ('nb09_golden_build', '골든셋 구성 → 평가 실행 (엑셀 연계)', '참고', CELLS.get('nb09'), '자동 평가셋',
  '검수 시트 · 확정 골든셋 · 실행 기록 양식', 'review · runrec', '평가팩 1 · 2 · 3'),
 ('nb10_answer_eval', '답변 품질 측정 — RAGAS 계열 지표', '핵심', CELLS.get('nb10'), '골든셋 · 청크 · LLM',
  '충실도 · 적합도 · 정확도 · 인용 · 거절 + 합격 판정', 'auto · summary', '이력대장 7_지표현황'),
 ('nb11_evidence_link', '질문에 근거 문서 연결 · 수집 공백 분리', '선택', CELLS.get('nb11'), '사외 60문항 · 청크',
  '질문별 근거 후보 · 판정 시트 · 수집 공백 목록', 'linked · gap', '평가팩 2d · 2e'),
 ('nb12_golden_llm', 'LLM 으로 골든셋 생성 · 자가검증 · 검수', '핵심', CELLS.get('nb12'), '청크 · LLM',
  '질문 · 정답 · 근거 + 자가검증 6항목 + 검수 시트', 'gen · chk · gold', '평가팩 1b_LLM골든셋_검수'),
 ('nb13_langgraph_rag', '노드 구성별 비교 — 개선 정량화', '핵심', CELLS.get('nb13'), '골든셋 · 청크 · LLM',
  '프리셋 v0~v6 지표 비교 · 증감 · 퇴행 문항', 'summ · dlt · detail', '이력대장 5 · 6'),
]
last = rows(ws, D3, cen=(3, 4), bold=(1,))
for i in range(5, last):
    if ws.cell(i, 3).value == '핵심':
        for j in range(1, 9):
            ws.cell(i, j).fill = OKF
    elif ws.cell(i, 3).value == '참고':
        for j in range(1, 9):
            ws.cell(i, j).fill = MUT

# ================= 4. 설정 기입 =================
ws = wb.create_sheet('4_설정_기입')
cols = ['구분', '항목', '무엇', '예시 · 기본값', '사내 값 (기입)', '확인']
head(ws, '4. 설정 기입표', '여기 적은 뒤 .env 와 lab_config.py 에 옮겨 넣기 · 비밀번호는 적지 말 것', cols,
     [12, 22, 44, 34, 34, 8])
D4 = [
 ('.env', 'LAB_MODE', '실행 모드', 'sample → live', '', ''),
 ('.env', 'SAMPLE_SET', '샘플 규모 (연습용)', '16 또는 100', '', ''),
 ('.env', 'SAMPLE_MODELS', '샘플 + 실제 모델 (노트북 연습용)', 'true / false', '', ''),
 ('.env', 'ORA_USER / ORA_DSN', 'Oracle 접속 — 읽기 전용 계정', 'host:1521/SERVICE', '', ''),
 ('.env', 'MILVUS_URI / COLLECTION', '벡터 DB 주소 · 컬렉션', 'http://host:19530', '', ''),
 ('.env', 'EMBED_MODE / EMBED_URL', '임베딩 방식 · 엔드포인트', 'local / api', '', ''),
 ('.env', 'RERANK_MODE / RERANK_URL', '리랭커 방식 · 엔드포인트', 'local / api / none', '', ''),
 ('.env', 'LLM_URL / LLM_MODEL', '사내 LLM (OpenAI 호환)', '.../v1/chat/completions', '', ''),
 ('lab_config', "RAW['table']", '통합 문서 테이블 이름', 'RAG_RAW_DOC', '', ''),
 ('lab_config', "RAW 컬럼 9개", 'doc_id · doc_type · title · body · published_at · collected_at · org_name · author 등', '위 이름 그대로', '', ''),
 ('lab_config', "CHUNK['table']", '청크 테이블 이름', 'RAG_CHUNK', '', ''),
 ('lab_config', 'CHUNK 컬럼 4개', 'chunk_id · doc_id · seq · text', '위 이름 그대로', '', ''),
 ('lab_config', 'MV 필드', 'Milvus 필드 이름 (pk · doc_id · text · dense · sparse · doc_type)', '없으면 None', '', ''),
 ('lab_config', 'DOC_TYPE_MAP', '사내 유형 코드 → 표준 7종 매핑', 'NEWS · BROKER · INSTITUTION · EMAIL · MEETING · REPORT · EXEC_REPORT', '', ''),
 ('lab_config', 'EMBED_MAX_LENGTH', '임베딩 최대 토큰 (현재 설정값 그대로)', '8192', '', ''),
]
last = rows(ws, D4, cen=(1, 6), inp=(5, 6), bold=(2,))
dv(ws, ['확인', '미정', '해당없음'], f'F5:F{last - 1}')
ws.cell(last + 1, 1, '주의').font = f(bold=True, color='5C440C')
x = ws.cell(last + 1, 2, '비밀번호 · API 키는 이 표에 적지 않는다. .env 에만 직접 입력하고 파일은 공유하지 않는다.')
x.font = f(color='5C440C'); x.fill = WARN; x.border = BOX; x.alignment = WRAP
ws.merge_cells(start_row=last + 1, start_column=2, end_row=last + 1, end_column=6)

# ================= 5. 안 돌리는 파일 =================
ws = wb.create_sheet('5_안돌리는_파일')
cols = ['파일', '역할', '직접 여는가']
head(ws, '5. 라이브러리 파일 13개 — 직접 실행하지 않음', '노트북이 불러 쓰는 코드. lab_config 만 예외적으로 수정', cols, [24, 76, 20])
D5 = [
 ('lab_config.py', '설정 단일 지점 — 접속 정보 · 테이블 · 컬럼 매핑 · 실행 모드', '★ 수정함'),
 ('lab_io.py', 'Oracle 읽기 / 샘플 제공 / 결과 저장', '안 열어도 됨'),
 ('lab_text.py', '소스별 정제 · 토큰 수 · 품질 지표 · 근사 중복', '안 열어도 됨'),
 ('lab_search.py', 'BM25 · 임베딩 · 리랭커 · Milvus · RRF · MMR · LLM 답변', '안 열어도 됨'),
 ('lab_pipeline.py', '적재 무결성 점검 기준(CRITERIA) · 소스별 필수 메타', '기준 조정 시'),
 ('lab_eval.py', '검색 지표 계산 — Hit · Recall · MRR · nDCG', '안 열어도 됨'),
 ('lab_ragas.py', '표준 지표 16종 정의 · 계산 · LLM 채점 프롬프트', '기준 조정 시'),
 ('lab_autoeval.py', '라벨 없는 평가셋 생성 · 근거 문서 연결', '안 열어도 됨'),
 ('lab_golden_llm.py', 'LLM 골든셋 생성 · 자가검증 · 유형 배분', '프롬프트 조정 시'),
 ('lab_graph.py', 'LangGraph 노드 · 프리셋 v0~v6 · 비교', '프리셋 추가 시'),
 ('lab_sql.py', '정형 연계 — 라우팅 · SQL 생성 · 정적 검증 · 실행', '허용목록 조정 시'),
 ('lab_sample.py', '샘플 데이터 16건 (sample 모드 기본)', '안 열어도 됨'),
 ('lab_sample100.py', '확장 샘플 100건 (SAMPLE_SET=100)', '안 열어도 됨'),
]
last = rows(ws, D5, cen=(3,), bold=(1,))
for i in range(5, last):
    if str(ws.cell(i, 3).value).startswith('★'):
        for j in range(1, 4):
            ws.cell(i, j).fill = WARN

from openpyxl.workbook.properties import CalcProperties      # noqa: E402
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.active = 0
wb.save(OUT)
print('saved', os.path.abspath(OUT), [w.title for w in wb.worksheets])
