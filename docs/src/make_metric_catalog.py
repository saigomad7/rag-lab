# -*- coding: utf-8 -*-
"""S22 정형 자산 카탈로그 양식 → metric_catalog_template_rev1.xlsx
   이 파일을 채우면 rag_lab/nb07_sql_router.py 가 그대로 읽어 프롬프트 · allow-list · 라우팅에 쓴다."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

OUT = 'metric_catalog_template_rev1.xlsx'
FONT = '맑은 고딕'
PRI = '185463'
thin = Side(style='thin', color='BFBFBF')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
F = lambda c: PatternFill('solid', start_color=c, end_color=c)
HDR, EX, INP = F('E7EDEE'), F('F2F7F7'), F('FFFDE7')
f = lambda **k: Font(name=FONT, size=k.pop('size', 10), **k)
WRAP = Alignment(wrap_text=True, vertical='top')
CEN = Alignment(horizontal='center', vertical='center', wrap_text=True)

SHEETS = [
 ('지표정의', '업무 용어를 SQL 식으로 1:1 정의한다 — Phase 3의 첫 작업이자 Text-to-SQL 의 근거',
  [('지표ID', 10), ('지표명 (업무 용어)', 18), ('별칭 (; 구분)', 22), ('한 줄 정의', 34), ('대상 뷰 · 테이블', 20),
   ('집계식 (SQL)', 30), ('기본 필터', 22), ('단위', 9), ('기간 단위', 10), ('소유 부서', 12), ('갱신 주기', 10), ('비고', 18)],
  [['M001', '출하량', '출하;shipment;판매량', '기간 중 고객에게 출하한 물량', 'V_SALES_MONTHLY',
    'SUM(SHIP_QTY_EB)', "STATUS='CONFIRMED'", 'EB', '월', '영업전략팀', '일 1회', '반품 제외'],
   ['M002', '계약가격', '계약가;contract price;ASP', '월 단위 확정 계약 단가(가중평균)', 'V_PRICE_MONTHLY',
    'SUM(AMT_USD)/NULLIF(SUM(QTY),0)', "PRICE_TYPE='CONTRACT'", 'USD/GB', '월', '마케팅팀', '월 1회', '현물가는 M003'],
   ['M003', '재고주수', '재고 주수;weeks of inventory;WOI', '기말 재고를 주 평균 출하량으로 나눈 값', 'V_INVENTORY_WEEKLY',
    'SUM(INV_QTY)/NULLIF(AVG(WEEKLY_SHIP),0)', '', '주', '주', 'SCM팀', '주 1회', '']]),
 ('코드값사전', '질문에 나오는 말("HBM3E 12단")을 코드로 바꾸는 표 — 라우팅과 WHERE 절에 쓰인다',
  [('구분', 10), ('코드', 14), ('표준명', 18), ('별칭 (; 구분)', 30), ('사용 컬럼', 20), ('비고', 20)],
  [['제품', 'P-HBM3E-12', 'HBM3E 12단', 'HBM3E 12H;HBM3E 12단;12단 HBM3E', 'PRODUCT_CODE', ''],
   ['제품', 'P-ESSD', 'eSSD', 'eSSD;enterprise SSD;기업용 SSD', 'PRODUCT_CODE', ''],
   ['고객', 'C-0001', '고객사 A', '고객사 A;A사', 'CUSTOMER_CODE', '실명은 코드로만 관리'],
   ['지역', 'R-CN', '중국', '중국;China;CN', 'REGION_CODE', '']]),
 ('테이블카탈로그', '연계 대상 테이블 · 뷰 목록 — 1차는 5~10개로 한정한다',
  [('물리명', 22), ('논리명', 18), ('유형', 9), ('소유 부서', 12), ('갱신 주기', 10), ('행 수 (대략)', 12),
   ('보존 기간', 10), ('PK', 18), ('주요 컬럼', 34), ('설명', 28), ('연계 우선', 9)],
  [['V_SALES_MONTHLY', '월별 출하', '뷰', '영업전략팀', '일 1회', '1,200,000', '5년', 'YM+PRODUCT_CODE+CUSTOMER_CODE',
    'YM;PRODUCT_CODE;CUSTOMER_CODE;REGION_CODE;SHIP_QTY_EB;AMT_USD', '확정 출하 실적(월 집계)', 'P0'],
   ['V_PRICE_MONTHLY', '월별 가격', '뷰', '마케팅팀', '월 1회', '80,000', '5년', 'YM+PRODUCT_CODE+PRICE_TYPE',
    'YM;PRODUCT_CODE;PRICE_TYPE;QTY;AMT_USD', '계약가 · 현물가', 'P0'],
   ['V_INVENTORY_WEEKLY', '주별 재고', '뷰', 'SCM팀', '주 1회', '300,000', '3년', 'YW+PRODUCT_CODE',
    'YW;PRODUCT_CODE;INV_QTY;WEEKLY_SHIP', '기말 재고 · 주 평균 출하', 'P1']]),
 ('allow_list', 'LLM 이 조회해도 되는 뷰와 컬럼 — 이 목록 밖의 SQL 은 실행 전에 차단된다',
  [('뷰 · 테이블', 24), ('허용 컬럼 (; 구분, * 면 전체)', 44), ('행 접근 제한 컬럼', 18), ('사용', 8), ('비고', 24)],
  [['V_SALES_MONTHLY', 'YM;PRODUCT_CODE;CUSTOMER_CODE;REGION_CODE;SHIP_QTY_EB;AMT_USD', 'REGION_CODE', 'Y', '고객 실명 컬럼 제외'],
   ['V_PRICE_MONTHLY', 'YM;PRODUCT_CODE;PRICE_TYPE;QTY;AMT_USD', '', 'Y', ''],
   ['V_INVENTORY_WEEKLY', '*', '', 'Y', ''],
   ['T_SALES_RAW', '', '', 'N', '원본 테이블 — 직접 조회 금지(S23)']]),
]

wb = Workbook()
ws = wb.active
ws.title = '안내'
ws['A1'] = 'S22 정형 자산 카탈로그 — 작성 양식 rev.1'
ws['A1'].font = f(size=14, bold=True, color=PRI)
ws['A2'] = '2026-09-20 · Phase 3(정형 데이터 연계)의 첫 작업 · 채운 뒤 rag_lab/nb07_sql_router.py 에서 그대로 읽는다'
ws['A2'].font = f(size=9, color='6B6B6B')
rows = [
 ('왜 필요한가', '업무 용어("수급 충족률", "계약가")를 SQL 식으로 못 적으면 Text-to-SQL 은 시작할 수 없다. 이 파일이 프롬프트 · allow-list · 라우팅의 근거가 된다.'),
 ('작성 순서', '① 테이블카탈로그(대상 좁히기) → ② allow_list(허용 컬럼) → ③ 지표정의(집계식) → ④ 코드값사전(별칭)'),
 ('회색 행', '예시입니다. 지우고 실제 값을 채웁니다. 연노랑 칸이 입력 칸입니다.'),
 ('별칭', '질문에 나올 법한 표현을 ; 로 나열합니다. 라우팅과 코드 변환에 그대로 쓰입니다.'),
 ('집계식', 'SELECT 절에 그대로 들어갈 수 있는 식으로 적습니다. 예: SUM(SHIP_QTY_EB), SUM(AMT_USD)/NULLIF(SUM(QTY),0)'),
 ('1차 범위', '테이블은 5~10개로 한정합니다. 전부 열지 않습니다(S23).'),
 ('파일 위치', 'rag_lab/catalog/metric_catalog.xlsx 로 저장하면 nb07 이 자동으로 읽습니다.'),
 ('주의', '고객 실명 · 개인정보 컬럼은 allow_list 에서 제외합니다. 원본 테이블은 사용=N 으로 둡니다.'),
]
ws['A4'] = '항목'; ws['B4'] = '내용'
for c in ('A4', 'B4'):
    ws[c].font = f(bold=True); ws[c].fill = HDR; ws[c].border = BOX; ws[c].alignment = CEN
for i, (a, b) in enumerate(rows, 5):
    ws.cell(i, 1, a).font = f(bold=True); ws.cell(i, 2, b).font = f()
    for c in (1, 2):
        ws.cell(i, c).border = BOX; ws.cell(i, c).alignment = WRAP
ws.column_dimensions['A'].width = 16; ws.column_dimensions['B'].width = 110

for name, desc, cols, examples in SHEETS:
    s = wb.create_sheet(name)
    s['A1'] = name; s['A1'].font = f(size=13, bold=True, color=PRI)
    s['A2'] = desc; s['A2'].font = f(size=9, color='6B6B6B')
    s.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    s.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(cols))
    for j, (c, w) in enumerate(cols, 1):
        x = s.cell(4, j, c); x.font = f(bold=True); x.fill = HDR; x.border = BOX; x.alignment = CEN
        s.column_dimensions[L(j)].width = w
    for i, ex in enumerate(examples, 5):
        for j, v in enumerate(ex, 1):
            x = s.cell(i, j, v); x.font = f(color='808080', italic=True); x.fill = EX; x.border = BOX; x.alignment = WRAP
    for i in range(5 + len(examples), 5 + len(examples) + 30):      # 입력용 빈 행
        for j in range(1, len(cols) + 1):
            x = s.cell(i, j); x.fill = INP; x.border = BOX; x.alignment = WRAP; x.font = f()
    last = 5 + len(examples) + 29
    s.freeze_panes = 'A5'
    s.auto_filter.ref = f'A4:{L(len(cols))}{last}'
    heads = [c for c, _ in cols]
    def dv(col_name, items):
        if col_name in heads:
            j = heads.index(col_name) + 1
            d = DataValidation(type='list', formula1='"' + ','.join(items) + '"', allow_blank=True)
            s.add_data_validation(d); d.add(f'{L(j)}5:{L(j)}{last}')
    dv('단위', ['EB', 'TB', 'GB', 'USD', 'USD/GB', 'KRW', '%', '주', '개'])
    dv('기간 단위', ['일', '주', '월', '분기', '연'])
    dv('갱신 주기', ['실시간', '일 1회', '주 1회', '월 1회', '분기 1회'])
    dv('유형', ['테이블', '뷰', '마트'])
    dv('구분', ['제품', '고객', '지역', '조직', '기타'])
    dv('연계 우선', ['P0', 'P1', 'P2'])
    dv('사용', ['Y', 'N'])
    s.page_setup.orientation = 'landscape'
    s.sheet_properties.pageSetUpPr.fitToPage = True
    s.page_setup.fitToWidth = 1; s.page_setup.fitToHeight = 0

wb.save(OUT)
print('saved', OUT, [w.title for w in wb.worksheets])
